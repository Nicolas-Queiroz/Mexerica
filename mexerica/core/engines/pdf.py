"""PDF to office, image and text formats with PyMuPDF, pdf2docx and python-pptx."""

import contextlib
import io
import logging
import re
from pathlib import Path

import pymupdf

from ..context import Context, OperationError
from . import documents, office

# MuPDF and the layout parser print hints on stdout; route them to logging instead.
pymupdf.set_messages(pylogging=True)

CONVERSIONS = {"pdf": ["docx", "odt", "pptx", "xlsx", "jpg", "png", "epub", "txt", "md", "html"]}

IMAGE_DPI = 300
SLIDE_DPI = 150


def open_pdf(source: Path) -> pymupdf.Document:
    try:
        document = pymupdf.open(source)
    except Exception as error:
        raise OperationError(f"PDF inválido: {error}") from error
    if document.needs_pass:
        document.close()
        raise OperationError("O PDF está protegido por senha. Use a ferramenta Remover senha antes.")
    return document


def render_pages(ctx: Context, document: pymupdf.Document, target: str, stem: str, dpi: int) -> list[Path]:
    """One image per page, numbered when there is more than one page."""
    produced = []
    total = len(document)
    width = len(str(total))
    for index, page in enumerate(document):
        ctx.check_cancelled()
        pixmap = page.get_pixmap(dpi=dpi, alpha=False)
        name = f"{stem}.{target}" if total == 1 else f"{stem}-{index + 1:0{width}d}.{target}"
        out = ctx.workdir / name
        if target == "jpg":
            pixmap.pil_save(out, format="JPEG", quality=90, optimize=True)
        else:
            pixmap.save(out)
        produced.append(out)
        ctx.progress((index + 1) / total)
    return produced


def to_docx(ctx: Context, source: Path) -> Path:
    from pdf2docx import Converter

    logging.getLogger("pdf2docx").setLevel(logging.WARNING)
    logging.getLogger().setLevel(logging.WARNING)
    out = ctx.workdir / f"{source.stem}.docx"
    converter = Converter(str(source))
    try:
        converter.convert(str(out))
    except Exception as error:
        raise OperationError(f"Não foi possível converter o PDF em DOCX: {error}") from error
    finally:
        converter.close()
    return out


def to_pptx(ctx: Context, document: pymupdf.Document, stem: str) -> Path:
    """Each page becomes a full-bleed picture slide sized like the page."""
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation()
    first = document[0].rect
    presentation.slide_width = Emu(int(first.width * 12700))
    presentation.slide_height = Emu(int(first.height * 12700))
    blank = presentation.slide_layouts[6]
    folder = ctx.subdir("slides")
    for index, page in enumerate(document):
        ctx.check_cancelled()
        image = folder / f"{index}.png"
        page.get_pixmap(dpi=SLIDE_DPI, alpha=False).save(image)
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image), 0, 0, presentation.slide_width, presentation.slide_height
        )
        ctx.progress((index + 1) / len(document))
    out = ctx.workdir / f"{stem}.pptx"
    presentation.save(out)
    return out


_NUMBER_BR = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+,\d+$")
_NUMBER_EN = re.compile(r"^-?\d+(\.\d+)?$")


def as_cell(value):
    """Turns numeric-looking text into numbers, accepting 1.234,56 as well as 1234.56."""
    if value is None:
        return None
    text = str(value).strip()
    if _NUMBER_EN.match(text):
        return float(text) if "." in text else int(text)
    if _NUMBER_BR.match(text):
        number = text.replace(".", "").replace(",", ".")
        return float(number) if "." in number else int(number)
    return text


def to_xlsx(ctx: Context, document: pymupdf.Document, stem: str) -> Path:
    """Detected tables become sheets; a PDF without tables gets its text lines, one per row."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, page in enumerate(document):
        ctx.check_cancelled()
        tables = page.find_tables().tables
        for number, table in enumerate(tables, start=1):
            title = f"Pág {index + 1}" + (f" ({number})" if len(tables) > 1 else "")
            sheet = workbook.create_sheet(title[:31])
            # Reading each cell's box keeps "12,50" whole; table.extract() splits some glyphs.
            for row in table.rows:
                sheet.append(
                    [as_cell(page.get_textbox(cell).strip()) if cell else None for cell in row.cells]
                )
        ctx.progress((index + 1) / len(document))
    if not workbook.worksheets:
        sheet = workbook.create_sheet("Texto")
        for page in document:
            for line in page.get_text().splitlines():
                if line.strip():
                    sheet.append([as_cell(part) for part in re.split(r"\s{2,}|\t", line.strip())])
    out = ctx.workdir / f"{stem}.xlsx"
    workbook.save(out)
    return out


def to_markdown(document: pymupdf.Document) -> str:
    try:
        import pymupdf4llm

        # The layout parser prints OCR hints on stdout; they mean nothing to the user.
        with contextlib.redirect_stdout(io.StringIO()):
            return pymupdf4llm.to_markdown(document, show_progress=False)
    except Exception:
        return "\n\n".join(page.get_text() for page in document)


def convert(source: Path, source_key: str, target: str, ctx: Context) -> list[Path]:
    stem = source.stem
    if target == "docx":
        return [to_docx(ctx, source)]
    if target == "odt":
        return [office.lo_convert_one(ctx, to_docx(ctx, source), "odt")]
    with open_pdf(source) as document:
        if target in ("jpg", "png"):
            return render_pages(ctx, document, target, stem, IMAGE_DPI)
        if target == "pptx":
            return [to_pptx(ctx, document, stem)]
        if target == "xlsx":
            return [to_xlsx(ctx, document, stem)]
        if target == "txt":
            out = ctx.workdir / f"{stem}.txt"
            out.write_text("\n\n".join(page.get_text() for page in document), encoding="utf-8")
            return [out]
        markdown = to_markdown(document)
    ctx.progress(0.6)
    md = ctx.workdir / f"{stem}.md"
    md.write_text(markdown, encoding="utf-8")
    if target == "md":
        return [md]
    return documents.pandoc(ctx, md, "markdown", target, stem)
