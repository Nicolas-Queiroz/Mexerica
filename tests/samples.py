"""Builds one small sample file per supported source format."""

import gzip
import shutil
import tarfile
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw

from mexerica.core.context import Context
from mexerica.core.engines import images, media, office

MARKER = "Mexerica"
# Formats nobody can produce on Linux without the original apps.
UNAVAILABLE = {"pages", "key", "numbers", "rar"}


def _image(path: Path) -> Image.Image:
    image = Image.new("RGBA", (320, 200), (255, 140, 0, 255))
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, 200, 160), fill=(30, 90, 200, 180))
    draw.text((60, 90), MARKER, fill="white")
    return image


def _pdf(path: Path) -> None:
    document = pymupdf.open()
    for number in (1, 2):
        page = document.new_page(width=595, height=842)
        page.insert_text((72, 72), f"{MARKER} página {number}", fontsize=18)
        if number == 1:
            rows = [["Produto", "Qtd", "Valor"], ["Laranja", "10", "12,50"], ["Limão", "4", "3,20"]]
            for row_index, row in enumerate(rows):
                for col_index, cell in enumerate(row):
                    rect = pymupdf.Rect(72 + col_index * 120, 120 + row_index * 24, 192 + col_index * 120, 144 + row_index * 24)
                    page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                    page.insert_text((rect.x0 + 4, rect.y1 - 7), cell, fontsize=11)
    document.save(path)


def _docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading(f"{MARKER} documento", level=1)
    document.add_paragraph("Parágrafo com acentuação: ação, maçã, você.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    document.save(path)


def _pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    for number in (1, 2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"{MARKER} slide {number}"
        slide.placeholders[1].text = "Conteúdo do slide"
    presentation.save(path)


def _xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendas"
    sheet.append(["Produto", "Qtd", "Preço", "Total"])
    sheet.append(["Laranja", 10, 1.25, "=B2*C2"])
    sheet.append([MARKER, 4, 0.8, "=B3*C3"])
    other = workbook.create_sheet("Resumo")
    other.append(["Total", "=SUM(Vendas!D2:D3)"])
    workbook.save(path)


def build(folder: Path) -> dict[str, Path]:
    """Creates the samples in folder and returns {format key: path}."""
    folder.mkdir(parents=True, exist_ok=True)
    work = folder / "_work"
    work.mkdir(exist_ok=True)
    ctx = Context(workdir=work)
    samples: dict[str, Path] = {}

    base = _image(folder)
    for key in images.RASTER:
        path = folder / f"imagem.{key}"
        images.save(base, path, key)
        samples[key] = path
    svg = folder / "vetor.svg"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80">'
        '<rect width="120" height="80" fill="orange"/><circle cx="60" cy="40" r="30" fill="green"/></svg>'
    )
    samples["svg"] = svg

    tone = folder / "tom.wav"
    media.run_ffmpeg(ctx, ["-f", "lavfi", "-i", "sine=frequency=440:duration=2", "-ac", "2", tone])
    for key in media.AUDIO:
        path = folder / f"audio.{key}"
        if key == "wav":
            shutil.copy(tone, path)
        else:
            media.run_ffmpeg(ctx, ["-i", tone, *media.AUDIO_CODECS[key], path])
        samples[key] = path

    clip = folder / "clipe.mp4"
    media.run_ffmpeg(
        ctx,
        [
            "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=24",
            "-f", "lavfi", "-i", "sine=frequency=330:duration=2",
            "-shortest", *media.video_codec_args("mp4"), clip,
        ],
    )
    info = media.probe(clip)
    for key in media.VIDEO:
        path = folder / f"video.{key}"
        if key == "mp4":
            shutil.copy(clip, path)
        else:
            media.encode_video(ctx, clip, path, key, info)
        samples[key] = path

    _pdf(folder / "documento.pdf")
    samples["pdf"] = folder / "documento.pdf"

    _docx(folder / "texto.docx")
    samples["docx"] = folder / "texto.docx"
    for key, filter_name in (("doc", "doc:MS Word 97"), ("odt", "odt"), ("rtf", "rtf")):
        produced = office.lo_convert(ctx, samples["docx"], key, filter_name=filter_name)[0]
        samples[key] = shutil.move(produced, folder / f"texto.{key}")
    md = folder / "notas.md"
    md.write_text(f"# {MARKER}\n\nTexto em **Markdown** com *ênfase*.\n\n- item 1\n- item 2\n", encoding="utf-8")
    samples["md"] = md
    txt = folder / "notas.txt"
    txt.write_text(f"{MARKER}\nLinha dois do texto.\n", encoding="utf-8")
    samples["txt"] = txt
    import pypandoc

    epub = folder / "livro.epub"
    pypandoc.convert_file(str(md), "epub3", outputfile=str(epub), extra_args=["--metadata", "title=Livro"])
    samples["epub"] = epub

    _pptx(folder / "slides.pptx")
    samples["pptx"] = folder / "slides.pptx"
    for key, filter_name in (("ppt", "ppt:MS PowerPoint 97"), ("odp", "odp")):
        produced = office.lo_convert(ctx, samples["pptx"], key, filter_name=filter_name)[0]
        samples[key] = shutil.move(produced, folder / f"slides.{key}")

    _xlsx(folder / "planilha.xlsx")
    samples["xlsx"] = folder / "planilha.xlsx"
    for key, filter_name in (("xls", "xls:MS Excel 97"), ("ods", "ods")):
        produced = office.lo_convert(ctx, samples["xlsx"], key, filter_name=filter_name)[0]
        samples[key] = shutil.move(produced, folder / f"planilha.{key}")
    csv_path = folder / "tabela.csv"
    csv_path.write_text(f"nome;valor\n{MARKER};1.234,56\nLimão;2,5\n", encoding="utf-8")
    samples["csv"] = csv_path
    tsv_path = folder / "tabela.tsv"
    tsv_path.write_text(f"nome\tvalor\n{MARKER}\t10\n", encoding="utf-8")
    samples["tsv"] = tsv_path

    srt = folder / "legenda.srt"
    srt.write_text(f"1\n00:00:01,000 --> 00:00:02,500\n{MARKER} fala\n\n2\n00:00:03,000 --> 00:00:04,000\nSegunda linha\n", encoding="utf-8")
    samples["srt"] = srt
    vtt = folder / "legenda.vtt"
    vtt.write_text(f"WEBVTT\n\nNOTE comentário\n\n00:01.000 --> 00:02.500 align:start\n{MARKER} fala\n", encoding="utf-8")
    samples["vtt"] = vtt

    content = folder / "pasta"
    (content / "sub").mkdir(parents=True, exist_ok=True)
    (content / "leia.txt").write_text(MARKER)
    (content / "sub" / "dado.txt").write_text("123")
    with zipfile.ZipFile(folder / "pacote.zip", "w") as archive:
        archive.write(content / "leia.txt", "leia.txt")
        archive.write(content / "sub" / "dado.txt", "sub/dado.txt")
    samples["zip"] = folder / "pacote.zip"
    with tarfile.open(folder / "pacote.tar", "w") as archive:
        archive.add(content, "pasta")
    samples["tar"] = folder / "pacote.tar"
    with tarfile.open(folder / "pacote.tar.gz", "w:gz") as archive:
        archive.add(content, "pasta")
    samples["tgz"] = folder / "pacote.tar.gz"
    with gzip.open(folder / "leia.txt.gz", "wb") as handle:
        handle.write(MARKER.encode())
    samples["gz"] = folder / "leia.txt.gz"
    ctx.run(["7z", "a", "-y", folder / "pacote.7z", "pasta"], cwd=folder)
    samples["7z"] = folder / "pacote.7z"

    shutil.rmtree(work, ignore_errors=True)
    return samples
